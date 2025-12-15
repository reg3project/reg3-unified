"""
XLSX Writer - Writes extracted data to Excel spreadsheets

Creates XLSX files with 'prodotti' and 'sku' sheets matching
the expected reference format.
"""

import os
from pathlib import Path
from typing import Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class XLSXWriter:
    """Write extracted data to XLSX format."""

    # Column definitions for prodotti sheet (Italian names matching reference)
    PRODOTTI_COLUMNS = [
        'categoria_prodotto', 'nome_prodotto', 'nome_asta', 'pagina_catalogo',
        'tipo_layout', 'immagine_principale', 'codici_modelli', 'titolo_prodotto',
        'descrizione_prodotto',
    ]

    # Column definitions for SKU sheet (Italian names matching reference)
    SKU_COLUMNS = [
        'SKU', 'Nome Modello', 'Tensione di alimentazione di rete',
        'Motore elettrico', 'Potenza max', 'Coppia max',
        'Peso', 'Dimensioni (LxPxH)', 'Grado di protezione',
        'Temperatura ambiente di esercizio', 'Frequenza di utilizzo',
        'Velocità dell\'anta', 'Corsa max', 'Condensatore marcia',
        'Condensatore di spunto', 'Finecorsa', 'Dispositivo di sblocco',
        'Apparecchiatura elettronica', 'Forza max di spinta',
        'Rapporto di riduzione', 'Larghezza max anta', 'Peso max anta',
        'Pignone', 'Encoder', 'Regolazione della forza',
    ]

    # Internal to Italian column name mapping
    FIELD_TO_ITALIAN = {
        'sku_code': 'SKU',
        'model': 'Nome Modello',
        'voltage': 'Tensione di alimentazione di rete',
        'motor_type': 'Motore elettrico',
        'power': 'Potenza max',
        'torque': 'Coppia max',
        'weight': 'Peso',
        'dimensions': 'Dimensioni (LxPxH)',
        'ip_rating': 'Grado di protezione',
        'temperature': 'Temperatura ambiente di esercizio',
        'cycles_hour': 'Frequenza di utilizzo',
        'speed': 'Velocità dell\'anta',
        'max_stroke': 'Corsa max',
        'capacitor_run': 'Condensatore marcia',
        'capacitor_start': 'Condensatore di spunto',
        'limit_switch': 'Finecorsa',
        'release': 'Dispositivo di sblocco',
        'control_unit': 'Apparecchiatura elettronica',
        'max_force': 'Forza max di spinta',
        'gear_ratio': 'Rapporto di riduzione',
        'max_width': 'Larghezza max anta',
        'max_gate_weight': 'Peso max anta',
        'pinion': 'Pignone',
        'encoder': 'Encoder',
        'force_regulation': 'Regolazione della forza',
        'quantity': 'Quantità',
        'component_type': 'Tipo Componente',
    }

    def __init__(self):
        self.workbook = Workbook()

    def create_workbook(self, content: Dict, specs: List[Dict]) -> Workbook:
        """Create a workbook with prodotti and sku sheets."""
        # Remove default sheet
        if 'Sheet' in self.workbook.sheetnames:
            del self.workbook['Sheet']

        # Create sheets
        self._create_prodotti_sheet(content)
        self._create_sku_sheet(specs, content)

        return self.workbook

    def _create_prodotti_sheet(self, content: Dict) -> None:
        """Create the prodotti (products) sheet."""
        ws = self.workbook.create_sheet('prodotti')

        # Write headers (Italian names)
        for col, header in enumerate(self.PRODOTTI_COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='DDEEFF', end_color='DDEEFF', fill_type='solid')

        # Map internal names to Italian column names
        row_data = {
            'categoria_prodotto': content.get('category', ''),
            'nome_prodotto': content.get('product_name', ''),
            'nome_asta': content.get('model', ''),
            'pagina_catalogo': content.get('pages', ''),
            'tipo_layout': '',
            'immagine_principale': ', '.join(content.get('images', [])[:1]) if content.get('images') else '',
            'codici_modelli': '',
            'titolo_prodotto': content.get('product_name', ''),
            'descrizione_prodotto': content.get('description', ''),
        }

        for col, header in enumerate(self.PRODOTTI_COLUMNS, 1):
            ws.cell(row=2, column=col, value=row_data.get(header, ''))

        # Adjust column widths
        self._auto_adjust_columns(ws)

    def _create_sku_sheet(self, specs: List[Dict], content: Dict) -> None:
        """Create the SKU specifications sheet."""
        ws = self.workbook.create_sheet('sku')

        # Write headers (Italian names)
        for col, header in enumerate(self.SKU_COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='DDEEFF', end_color='DDEEFF', fill_type='solid')

        # Write data rows
        row_num = 2
        for spec in specs:
            # Add model from content if not in spec
            if 'model' not in spec:
                spec['model'] = content.get('product_name', '')

            for col, italian_name in enumerate(self.SKU_COLUMNS, 1):
                # Try to find value using Italian name directly, then try English mapping
                value = spec.get(italian_name, '')
                if not value:
                    # Look up English field name and get value
                    for eng_name, ita_name in self.FIELD_TO_ITALIAN.items():
                        if ita_name == italian_name:
                            value = spec.get(eng_name, '')
                            break
                ws.cell(row=row_num, column=col, value=value)
            row_num += 1

        # If no specs, write at least one row with model
        if not specs:
            ws.cell(row=2, column=2, value=content.get('product_name', ''))

        # Adjust column widths
        self._auto_adjust_columns(ws)

    def _auto_adjust_columns(self, ws) -> None:
        """Auto-adjust column widths based on content."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            # Set width with some padding
            adjusted_width = min(max_length + 2, 50)  # Max 50 chars wide
            ws.column_dimensions[column_letter].width = adjusted_width

    def save(self, output_path: str) -> None:
        """Save the workbook to file."""
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(output_path)


def write_xlsx(content: Dict, specs: List[Dict], output_path: str) -> str:
    """Write extracted data to XLSX file.

    Returns: Path to created file
    """
    writer = XLSXWriter()
    writer.create_workbook(content, specs)
    writer.save(output_path)
    return output_path
